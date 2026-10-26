"""Carrega o arquivo Excel exportado do sistema UFRGS."""
import re
import pandas as pd
import openpyxl
from pathlib import Path

_ABAS_IGNORADAS = {"naturezas", "Página6", "Leia-me", "Instruções", "Capa"}
_ETAPA_RE = re.compile(r'Etapa\s+(\d+)', re.IGNORECASE)
_CODIGO_RE = re.compile(r'\b([A-Z]{2,}[0-9]{2,}(?:/[0-9]+)*)\b')


def _parse_prereqs(texto: str) -> tuple[str, list[str]]:
    """Separa nome e pré-requisitos de uma célula combinada (com \n)."""
    if not texto or not isinstance(texto, str):
        return str(texto or ""), []
    partes = texto.split('\n', 1)
    nome = partes[0].strip()
    prereqs = []
    if len(partes) > 1:
        prereqs = _CODIGO_RE.findall(partes[1])
    return nome, prereqs


def load_excel(path: str | Path) -> dict[str, pd.DataFrame]:
    """Lê o arquivo Excel e retorna um dict {nome_aba: DataFrame}."""
    wb = openpyxl.load_workbook(path, data_only=True)
    resultado = {}

    for nome_aba in wb.sheetnames:
        if nome_aba in _ABAS_IGNORADAS:
            continue
        ws = wb[nome_aba]
        linhas = [row for row in ws.iter_rows(values_only=True)]
        if not linhas:
            continue

        # Detecta o header — busca linha com "Código" ou "Atividade"
        header_idx = 0
        for i, linha in enumerate(linhas[:5]):
            vals = [str(v or "").strip() for v in linha]
            if any("código" in v.lower() or "atividade" in v.lower() for v in vals):
                header_idx = i
                break

        header = [str(v or "").strip() for v in linhas[header_idx]]
        dados = []
        for linha in linhas[header_idx + 1:]:
            if not any(v for v in linha):
                continue
            row_dict = dict(zip(header, linha))
            dados.append(row_dict)

        if dados:
            resultado[nome_aba] = pd.DataFrame(dados)

    return resultado
